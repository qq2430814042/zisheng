# -*- coding: utf-8 -*-
"""md_to_docx.py — 把 docs/*.md 转成提交用 Word（与仓库同源，防止脚本丢失）。

用法: python tools/md_to_docx.py
"""
import glob
import os
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOCS = os.path.join(ROOT, "docs")

BODY_FONT = "宋体"
HEAD_FONT = "微软雅黑"


def set_font(run, name, size, bold=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_runs(p, text, name, size, bold=False):
    # 处理 **加粗** 内联
    for i, seg in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if not seg:
            continue
        run = p.add_run(seg)
        set_font(run, name, size, bold or (i % 2 == 1))


def flush_table(doc, rows):
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=ncols)
    t.style = "Table Grid"
    for ri, row in enumerate(rows):
        for ci in range(ncols):
            cell = t.cell(ri, ci)
            cell.text = ""
            p = cell.paragraphs[0]
            val = row[ci] if ci < len(row) else ""
            add_runs(p, val, BODY_FONT, 9.5, bold=(ri == 0))
    doc.add_paragraph()


def convert(md_path, docx_path):
    lines = open(md_path, encoding="utf-8").read().splitlines()
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = BODY_FONT
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)

    table_rows = []
    for raw in lines:
        line = raw.rstrip()
        stripped = line.strip()
        if stripped.startswith("|"):
            cells = [c.strip() for c in stripped.strip("|").split("|")]
            if all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                continue  # 分隔行
            table_rows.append(cells)
            continue
        if table_rows:
            flush_table(doc, table_rows)
            table_rows = []

        if not stripped:
            continue
        if stripped.startswith("### "):
            p = doc.add_paragraph()
            add_runs(p, stripped[4:], HEAD_FONT, 12, bold=True)
        elif stripped.startswith("## "):
            p = doc.add_paragraph()
            add_runs(p, stripped[3:], HEAD_FONT, 13.5, bold=True)
            p.space_before = Pt(10)
        elif stripped.startswith("# "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_runs(p, stripped[2:], HEAD_FONT, 17, bold=True)
        elif stripped.startswith("> "):
            p = doc.add_paragraph()
            add_runs(p, stripped[2:], BODY_FONT, 10)
            for r in p.runs:
                r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        elif re.match(r"^-\s*\[[ xX]\]\s*", stripped):
            p = doc.add_paragraph(style="List Bullet")
            mark = "☑ " if stripped[3] in "xX" else "☐ "
            add_runs(p, mark + re.sub(r"^-\s*\[[ xX]\]\s*", "", stripped), BODY_FONT, 10.5)
        elif stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, stripped[2:], BODY_FONT, 10.5)
        elif re.match(r"^\d+\.\s+", stripped):
            p = doc.add_paragraph(style="List Number")
            add_runs(p, re.sub(r"^\d+\.\s+", "", stripped), BODY_FONT, 10.5)
        else:
            p = doc.add_paragraph()
            add_runs(p, stripped, BODY_FONT, 10.5)
    if table_rows:
        flush_table(doc, table_rows)
    doc.save(docx_path)
    print("saved", docx_path)


def main():
    for md in sorted(glob.glob(os.path.join(DOCS, "*.md"))):
        convert(md, md[:-3] + ".docx")


if __name__ == "__main__":
    main()
