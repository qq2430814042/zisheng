# -*- coding: utf-8 -*-
"""build_ppt.py — 生成《字·生》数字体验课件 PPT（字生-演示课件.pptx）

深色墨金风、16:9、含讲稿备注；截图取自 assets/submission/（提交截图同款）。
用法: python tools/build_ppt.py
"""
import os
import glob
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SHOTS = os.path.join(ROOT, "assets", "submission")
OUT = os.path.join(ROOT, "字生-演示课件.pptx")

BG = "0B0B0F"
CARD = "15151C"
INK = "E9E6DF"
DIM = "A8A296"
GOLD = "E0B36A"
LINE = "34343E"
KAI = "楷体"
YAHEI = "微软雅黑"
W, H = 13.333, 7.5


def shot(name):
    hits = glob.glob(os.path.join(SHOTS, "*" + name + "*.jpg"))
    if not hits:
        raise SystemExit("缺少截图: " + name)
    return hits[0]


prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)


def cjk(run, name):
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", name)


def set_alpha(shape, pct):
    clr = shape._element.spPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
    clr.append(clr.makeelement(qn("a:alpha"), {"val": str(int(pct * 1000))}))


def rect(slide, x, y, w, h, color, alpha=None, line=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    if color:
        shp.fill.solid()
        shp.fill.fore_color.rgb = RGBColor.from_string(color)
        if alpha is not None:
            set_alpha(shp, alpha)
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = RGBColor.from_string(line)
        shp.line.width = Pt(0.9)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def frame(slide, x, y, w, h):
    rect(slide, x, y, w, h, None, line=LINE)


def tb(slide, x, y, w, h, blocks):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    first = True
    for b in blocks:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = b.get("align", PP_ALIGN.LEFT)
        if b.get("before"):
            p.space_before = Pt(b["before"])
        p.line_spacing = b.get("line", 1.18)
        for r in b["runs"]:
            run = p.add_run()
            run.text = r[0]
            run.font.size = Pt(r[1])
            run.font.bold = r[3] if len(r) > 3 else False
            run.font.color.rgb = RGBColor.from_string(r[2])
            cjk(run, r[4] if len(r) > 4 else YAHEI)
    return box


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def new_slide(img=None, overlay=None):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, W, H, BG)
    if img:
        pic = s.shapes.add_picture(img, 0, 0, width=Inches(W))
        if pic.height > Inches(H):
            pic.top = Inches((H - pic.height / 914400) / 2)
        if overlay:
            rect(s, 0, 0, W, H, "000000", alpha=overlay)
    return s


def head(s, kicker, title, sub=None):
    tb(s, 0.62, 0.42, 11.8, 0.4, [{"runs": [(kicker, 13, GOLD, False)]}])
    rect(s, 0.64, 0.82, 0.5, 0.028, GOLD)
    blocks = [{"runs": [(title, 30, INK, False, KAI)]}]
    if sub:
        blocks.append({"runs": [(sub, 13.5, DIM)], "before": 6})
    tb(s, 0.62, 0.98, 12.1, 1.0, blocks)


def foot(s, page):
    tb(s, 0.62, 7.03, 8, 0.3, [{"runs": [("《字·生》· 作品编号 316056 · 牛景飞（北京语言大学）", 10, "6B675E")]}])
    tb(s, 12.2, 7.03, 0.9, 0.3, [{"runs": [(page, 10, "6B675E")], "align": PP_ALIGN.RIGHT}])


def img_right(s, path, x=7.45, y=2.32, w=5.3):
    pic = s.shapes.add_picture(path, Inches(x), Inches(y), width=Inches(w))
    h = pic.height / 914400
    frame(s, x - 0.022, y - 0.022, w + 0.044, h + 0.044)
    return pic


def img_big(s, path, y=1.7, w=9.4):
    x = (W - w) / 2
    pic = s.shapes.add_picture(path, Inches(x), Inches(y), width=Inches(w))
    h = pic.height / 914400
    frame(s, x - 0.022, y - 0.022, w + 0.044, h + 0.044)
    return pic


def bullets(s, x, y, w, items, size=14):
    blocks = []
    for it in items:
        blocks.append({"runs": [("•  ", size, GOLD), (it, size, INK)], "before": 10, "line": 1.32})
    tb(s, x, y, w, 4.6, blocks)


# ---------- S1 封面 ----------
s = new_slide()
rows = ["日", "月", "山", "水", "人"]
for i, ch in enumerate(rows):
    tb(s, 10.6, 1.35 + i * 1.06, 2.2, 1.0, [{"runs": [(ch, 44, "26262E", False, KAI)], "align": PP_ALIGN.CENTER}])
tb(s, 0.8, 0.72, 9.2, 0.4, [{"runs": [("第十四届全国大学生数字媒体科技作品及创意竞赛 · 指定命题「科技+创意，讲好中国故事」", 12.5, GOLD)]}])
tb(s, 0.8, 2.35, 9.2, 1.6, [{"runs": [("字·生", 64, INK, False, KAI)]}])
tb(s, 0.8, 3.95, 9.2, 0.7, [{"runs": [("汉字五体演变的粒子交互体验 · 数字体验课件", 20, INK, False, KAI)]}])
rect(s, 0.85, 4.75, 2.2, 0.03, GOLD)
tb(s, 0.8, 5.05, 9.2, 1.3, [
    {"runs": [("适用：语文 / 文化 / 国际中文教育课堂（HSK 3 级及以上）· 40–45 分钟", 13, DIM)]},
    {"runs": [("作品编号 316056 ｜ 作者：牛景飞（北京语言大学）｜ 个人参赛 · 2026", 13, DIM)], "before": 8},
])
notes(s, "开场白：三千年前的『日』是什么样子？这堂课我们不用课本插图，而是亲手把汉字的三千年『推』一遍。作品是纯前端交互网页，双击即用、完全离线。")

# ---------- S2 课件概览 ----------
s = new_slide()
head(s, "课件说明", "一堂可以「动手」的汉字课")
cards = [
    ("01 · 适用对象", ["中学语文与文化课", "高校通识 / 美育课", "国际中文教育课堂（HSK 3+）"]),
    ("02 · 建议课时", ["共 40–45 分钟", "导入 5′ · 探究 10′", "实践 15′ · 检测 10′ · 总结 5′"]),
    ("03 · 使用方式", ["双击「字生-单文件版.html」", "完全离线 · 无需安装", "电脑 / 平板 / 手机均可"]),
]
for i, (t, lines) in enumerate(cards):
    x = 0.62 + i * 4.1
    rect(s, x, 2.35, 3.8, 2.35, CARD, line=LINE)
    rect(s, x, 2.35, 3.8, 0.045, GOLD)
    blocks = [{"runs": [(t, 15, GOLD, True)]}]
    for ln in lines:
        blocks.append({"runs": [(ln, 13, INK)], "before": 8, "line": 1.25})
    tb(s, x + 0.28, 2.62, 3.3, 2.0, blocks)
tb(s, 0.62, 5.05, 12.1, 1.6, [
    {"runs": [("教学目标　", 14, GOLD, True), ("知识：", 14, INK, True), ("汉字的五体演变与『六书』造字智慧　　", 14, INK), ("能力：", 14, INK, True), ("动手构字、书写与观察分析　　", 14, INK), ("价值：", 14, INK, True), ("感受汉字之美，建立文化认同", 14, INK)], "line": 1.4},
    {"runs": [("教学流程　", 14, GOLD, True), ("导入（观演变）→ 探究（一字千年）→ 实践（构字 / 写字）→ 检测（认字闯关）→ 总结（字量讨论），全部环节由同一件交互作品串联。", 14, INK)], "before": 10, "line": 1.4},
])
foot(s, "02 / 15")
notes(s, "说明课件的双身份：既是可上手的课堂教学工具，也是作品的完整展示。所有环节都在同一个网页里完成，不需要来回切换软件。可按课时拆分：只做导入+探究约 15 分钟；实践环节可单独成课。")

# ---------- S3 导入·五体演变 ----------
s = new_slide()
head(s, "环节一 · 导入（约 5 分钟）", "「日」在三千年前，长什么样？")
img_right(s, shot("02甲骨文"))
bullets(s, 0.62, 2.45, 6.6, [
    "滚轮逐幕观演：甲骨 → 金文 → 简帛 → 小篆 → 隶书 → 楷书",
    "观察提问：哪一幕的变化最大？为什么？",
    "点画面中的大字：看它的结构与本义释读",
], size=14)
foot(s, "03 / 15")
notes(s, "先让画面自己说话：不要急着讲解，让学生连续观演两轮，再抛出问题。提示留意两处转折：金文到简帛的『隶变』苗头、小篆到隶书的彻底转向。点字弹出释读卡可用于随机抽问。")

# ---------- S4 简帛一幕 ----------
s = new_slide()
head(s, "环节一 · 导入", "隶变：从「画出来」到「写顺手」", "简帛一幕——篆隶之间的手写墨迹：书写效率，推动了字体的演变")
img_big(s, shot("03简帛"), y=2.02, w=8.8)
foot(s, "04 / 15")
notes(s, "汉字演变的最大动力不是美学而是效率：竹简上写篆书太慢，于是笔画开始连写、拉直、简化，隶书由此萌芽。这一幕的墨迹感可以让学生直观看到『书写』的痕迹。")

# ---------- S5 探究·一字千年 ----------
s = new_slide()
head(s, "环节二 · 探究（约 10 分钟）", "任选一个字，亲手推动它的三千年")
img_right(s, shot("04一字千年"))
bullets(s, 0.62, 2.45, 6.6, [
    "全量字库 14,873 字 · 28,776 个历代字形",
    "拖时间轴：从它出现的最早年代，一路演化到楷书",
    "讨论任务：「我」「学」「家」三个字的来路有什么不同？",
], size=14)
foot(s, "05 / 15")
notes(s, "这是最核心的探究工具。让学生自己选字、自己拖轴，教师只在关键节点点拨：比如『学』字上面是两只手在教孩子结绳计数。任务化：每组挑一个字，向全班讲述它的演变故事。")

# ---------- S6 探究·检索与诚实标注 ----------
s = new_slide()
head(s, "环节二 · 探究", "找不到古字形？作品会诚实地说出来")
img_right(s, shot("05字库检索"))
bullets(s, 0.62, 2.45, 6.6, [
    "按 214 部首 / 36 笔画检索任意字，或「随机一字」",
    "有的字「出生」很晚：甲骨金文里本就还没有它",
    "缺哪一阶段就标注哪一阶段——这本身就是汉字的故事",
], size=14)
foot(s, "06 / 15")
notes(s, "借『晚出字』做一次科学精神教育：数据不会撒谎，工具要说真话。可以现场检索『妈』『茶』『灯』等生活字，看看它们最古老的样子出现在哪个年代。")

# ---------- S7 实践①·构字 ----------
s = new_slide()
head(s, "环节三 · 实践（约 15 分钟）", "日 + 月 = 明：亲手拼出会意字")
img_right(s, shot("06构字"))
bullets(s, 0.62, 2.45, 6.6, [
    "选字 / 手写双入口：8,600 组部件组合",
    "部首归一：手写「水」+「可」，也能合成「河」",
    "合出来的字可「看它三千年」；拆解释义可展开",
    "课堂活动：同桌各写一个部件，合起来猜是什么字",
], size=13.5)
foot(s, "07 / 15")
notes(s, "从『看』到『做』的关键一步。会意字最能体现先民的思维：日与月合为『明』。手写通道让学生必须真正写出部件，而不是点按钮；合成后立刻能看三千年，形成『造字—演变』的闭环。")

# ---------- S8 实践②·写字 ----------
s = new_slide()
head(s, "环节三 · 实践", "写下的每一笔，粒子都会「接住」", "三种模式：自由书写 · 跟着写（描红评分）· 名字图　｜　离线识别引擎 12,361 字（合成样本 Top-1 97% / Top-3 100%）")
img_big(s, shot("07手写识别"), y=2.02, w=8.8)
foot(s, "08 / 15")
notes(s, "写字环节强调『手写』的必要性：识别引擎会对 Top-3 候选让你确认，写歪了也没关系。『跟着写』模式有评分，适合低年级与留学生入门；写出的任何字都会化入粒子星河。")

# ---------- S9 实践③·名字图 ----------
s = new_slide()
head(s, "环节三 · 实践", "把名字变成一片星空", "输入 2–4 个字生成粒子星图，可保存为纪念海报（1920×1080）")
img_big(s, shot("08名字图"), y=2.02, w=8.8)
foot(s, "09 / 15")
notes(s, "情感化设计：把名字变成作品，学生可以把海报带回家。课堂用法——写同桌的名字，互赠『星空』；也可以写班级名、校名。")

# ---------- S10 检测·认字 ----------
s = new_slide()
head(s, "环节四 · 检测（约 10 分钟）", "三关闯关：认形 → 知义 → 辨体", "三档难度可选：简单 13 题（热身）· 普通 20 题（默认）· 困难 25 题（认形从全字库随机抽取、同部首干扰）")
img_right(s, shot("09认字"))
bullets(s, 0.62, 2.45, 6.6, [
    "认形：看甲骨文，猜今天的字；困难档同部首干扰",
    "知义：它的本义是什么？答对附一条汉字小知识",
    "辨体：两个字形，哪一个更早？困难档比相邻字体",
    "每轮随机抽题、跨轮不重样；实时总分",
], size=13.5)
foot(s, "10 / 15")
notes(s, "把检测做成游戏：可以个人挑战，也可以分组抢答。三档难度让不同基础的学生都有的玩：简单档热身、普通档日常、困难档从全字库随机出题（几乎每局都不一样，认形干扰项还同部首）。答错的即时讲解可以顺带复习前面的知识点。")

# ---------- S11 总结·字量 ----------
s = new_slide()
head(s, "环节五 · 总结（约 5 分钟）", "汉字为什么越来越多，又「稳」了下来？")
img_right(s, shot("10字量"))
bullets(s, 0.62, 2.45, 6.6, [
    "甲骨约 4,500（已释读约 1,500）→ 康熙字典 47,035 → 通用规范 8,105",
    "作品字库实录：14,873 字 · 28,776 个历代字形",
    "讨论：规范用字时代，我们学汉字学的是什么？",
], size=14)
foot(s, "11 / 15")
notes(s, "用数据收束全课：字量在增长，但『规范』让它稳定在 8,105 字。引导学生意识到：认识几千个常用字，就掌握了通向三千年的钥匙。")

# ---------- S12 课后任务 ----------
s = new_slide()
head(s, "拓展 · 课后任务", "去字墙上，找找自己名字里的「最古老汉字」")
tasks = [
    ("任务一", "在「尾声」字墙点三个字，记下它们各自最早的年代"),
    ("任务二", "用「构字」拼出一个你名字里的字，读一读它的拆解"),
    ("任务三", "在「名字图」写下同桌的名字，比一比谁的字更『古老』"),
]
for i, (t, txt) in enumerate(tasks):
    y = 2.5 + i * 1.15
    rect(s, 0.62, y, 12.1, 0.95, CARD, line=LINE)
    rect(s, 0.62, y, 0.045, 0.95, GOLD)
    tb(s, 1.0, y + 0.1, 2.2, 0.7, [{"runs": [(t, 16, GOLD, True, KAI)]}])
    tb(s, 3.0, y + 0.18, 9.4, 0.7, [{"runs": [(txt, 14.5, INK)]}])
tb(s, 0.62, 6.25, 12.1, 0.5, [{"runs": [("对应章节：尾声（字墙）· 构字 · 写字（名字图）", 12, DIM)]}])
foot(s, "12 / 15")
notes(s, "课后任务把课堂延伸到家庭：学生回家继续玩作品即可完成，无需额外材料。任务三适合作为下节课的开场分享。")

# ---------- S13 作品信息 ----------
s = new_slide()
head(s, "作品信息", "一个纯前端、完全离线的数字体验")
bullets(s, 0.62, 2.4, 7.6, [
    "单文件版 47.3 MB：字体、数据、识别模型全部内嵌，双击即用、断网可用",
    "数万枚实时粒子（Three.js + 自定义着色器）实时生成全部画面",
    "离线手写识别（WebAssembly）· 字体按需分块 · 字库分片懒加载",
    "支持桌面 / 平板 / 手机（320px 宽度起）；无障碍与动效降级适配",
], size=14)
rect(s, 8.6, 2.4, 4.12, 3.9, CARD, line=LINE)
rect(s, 8.6, 2.4, 4.12, 0.045, GOLD)
tb(s, 8.9, 2.7, 3.5, 3.4, [
    {"runs": [("体验方式", 15, GOLD, True)]},
    {"runs": [("双击「字生-单文件版.html」", 13, INK)], "before": 8, "line": 1.35},
    {"runs": [("主副本：Gitee（见提交链接）", 13, INK)], "before": 4, "line": 1.35},
    {"runs": [("源码镜像：GitHub", 13, INK)], "before": 4, "line": 1.35},
    {"runs": [("作品编号：316056", 13, INK)], "before": 14, "line": 1.35},
    {"runs": [("作者：牛景飞", 13, INK)], "before": 4, "line": 1.35},
    {"runs": [("单位：北京语言大学", 13, INK)], "before": 4, "line": 1.35},
])
foot(s, "13 / 15")
notes(s, "强调工程的完成度：不是演示原型，而是可以直接交付使用的成品——单文件双击即用，评委与学生都不需要配置环境。")

# ---------- S14 素材与版权 ----------
s = new_slide()
head(s, "素材与版权", "零第三方图片与音频素材")
bullets(s, 0.62, 2.4, 12.1, [
    "全部画面与音效由程序实时生成（五声音阶 · 四种可选音色：古琴 / 编钟 / 洞箫 / 水滴）",
    "开源致谢：Three.js（MIT）· LXGW WenKai 楷体（SIL OFL 1.1）· HanScribe 手写识别（MIT）· character-evolution-dataset-1bit 古文字字形集（MIT）· CCD 汉字拆解表（Wikimedia Commons）",
    "AI 使用说明：AI 仅用于代码辅助与文字整理；画面、音效、字形数据均非 AI 生成",
    "字体与数据的许可声明随作品与源码一并保留；CCD 拆解数据在界面中署名",
], size=14)
foot(s, "14 / 15")
notes(s, "版权页是竞赛材料的重要组成：所有素材来源可追溯、许可合规。若评委问到数据来源，逐一对应：字形来自 MIT 数据集，拆解数据来自 Wikimedia Commons 项目页并已署名。")

# ---------- S15 结束 ----------
s = new_slide()
tb(s, 0, 2.5, W, 1.6, [{"runs": [("字·生", 60, INK, False, KAI)], "align": PP_ALIGN.CENTER}])
rect(s, 5.57, 4.15, 2.2, 0.03, GOLD)
tb(s, 0, 4.45, W, 0.6, [{"runs": [("让沉默的字符，在粒子系统中重新「活」起来", 18, DIM)], "align": PP_ALIGN.CENTER}])
tb(s, 0, 5.6, W, 0.5, [{"runs": [("谢谢观看 · 欢迎体验　｜　作品编号 316056", 12.5, "6B675E")], "align": PP_ALIGN.CENTER}])
notes(s, "结束页用于课堂收束与展示收尾：邀请听众亲自打开单文件版，从『日』开始，走一遍他们的三千年。")

prs.core_properties.title = "字·生 — 汉字五体演变的粒子交互体验（数字体验课件）"
prs.core_properties.author = "牛景飞"
prs.core_properties.subject = "第十四届全国大学生数字媒体科技作品及创意竞赛 · 作品编号 316056"
prs.core_properties.comments = "由 tools/build_ppt.py 生成；截图与文案与单文件版作品同步。"

prs.save(OUT)
print("written:", OUT, os.path.getsize(OUT), "bytes,", len(prs.slides._sldIdLst), "slides")
